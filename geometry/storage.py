# -*- coding: utf-8 -*-
"""Хранение результатов расчётов: sqlite3 + экспорт в docx/xlsx."""

import sqlite3
from datetime import datetime

from docx import Document
from openpyxl import Workbook


class ResultsDB:
    """Обёртка над sqlite3 для таблицы результатов."""

    def __init__(self, path="results.db"):
        self._conn = sqlite3.connect(path)
        self._conn.execute(
            "CREATE TABLE IF NOT EXISTS results ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, "
            "body TEXT, params TEXT, material TEXT, "
            "volume REAL, area REAL, mass REAL, created TEXT)"
        )

    def add(self, body):
        """Сохранить результат расчёта тела."""
        cur = self._conn.execute(
            "INSERT INTO results (body, params, material, volume, area, mass, created) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (body.kind(), body.params_str(), str(body.material),
             body.volume(), body.surface_area(), body.mass,
             datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        )
        self._conn.commit()
        return cur.lastrowid

    def all(self):
        """Все записи в виде списка кортежей."""
        return self._conn.execute(
            "SELECT id, body, params, material, volume, area, mass, created "
            "FROM results ORDER BY id"
        ).fetchall()

    def clear(self):
        """Очистить таблицу."""
        self._conn.execute("DELETE FROM results")
        self._conn.commit()

    def __len__(self):
        return self._conn.execute("SELECT COUNT(*) FROM results").fetchone()[0]

    def __iter__(self):
        return iter(self.all())


def export_docx(rows, filename="report.docx"):
    """Выгрузить таблицу результатов в docx."""
    doc = Document()
    doc.add_heading("Результаты расчёта геометрических тел", level=1)
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ("ID", "Тело", "Параметры", "Материал", "V, м^3", "S, м^2", "m, кг", "Дата")
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = f"{value:.4f}" if isinstance(value, float) else str(value)
    doc.save(filename)
    return filename


def export_xlsx(rows, filename="report.xlsx"):
    """Выгрузить таблицу результатов в xlsx."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    headers = ("ID", "Тело", "Параметры", "Материал", "V, м^3", "S, м^2", "m, кг", "Дата")
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col, width in zip("ABCDEFGH", (6, 16, 20, 26, 12, 12, 12, 20)):
        ws.column_dimensions[col].width = width
    wb.save(filename)
    return filename
